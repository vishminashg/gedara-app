"""Builds IT3060_Milestone03_Usability_Workbook.xlsx: an EMPTY results workbook for the working-app usability test.

Run:  python3 make_workbook.py [output.xlsx]
Input cells (yellow) are left blank on purpose. No participant, result or issue is ever pre-filled.
Formulas return "" on blank input so the sheets show no errors until real data is typed in.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
import data3 as D

NAVY, YEL, SOFT, BLUE = "184F95", "FFF4C2", "F6F6F4", "E3EEF9"
hf = Font(name="Arial", bold=True, color="FFFFFF", size=10)
hfill = PatternFill("solid", fgColor=NAVY)
inp = PatternFill("solid", fgColor=YEL)
soft = PatternFill("solid", fgColor=SOFT)
tgt_fill = PatternFill("solid", fgColor=BLUE)
thin = Side(style="thin", color="D9D9D3")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name="Arial", size=10)
bold = Font(name="Arial", size=10, bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
green = PatternFill("solid", bgColor="D6EFE0")
red = PatternFill("solid", bgColor="F8D7D3")
amber = PatternFill("solid", bgColor="FDE9C8")

NP = 10          # participant slots U1..U10
NI = 30          # issue rows
LOG_LAST = 301   # Summary formulas read Task log rows 2..301
FIRST_FREE = NP * len(D.TASKS) + 2   # first free-entry row of the Task log (after the U x T grid)


def header(ws, row, cols, height=None):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = hf; c.fill = hfill; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if height:
        ws.row_dimensions[row].height = height


def widths(ws, ws_):
    for i, w in enumerate(ws_, 1):
        ws.column_dimensions[L(i)].width = w


def inputs(ws, row, c1, c2, wrap_text=False):
    for cidx in range(c1, c2 + 1):
        c = ws.cell(row=row, column=cidx)
        c.fill = inp; c.border = box; c.font = base
        if wrap_text:
            c.alignment = wrap


def build(out, hook=None):
    wb = Workbook()

    # ------------------------------------------------------------ Read me
    ws = wb.active; ws.title = "Read me"
    lines = [
        ("IT3060 Milestone 03 – Gedara usability test results workbook (working app)", "title"),
        ("This workbook is EMPTY. Every number in it must come from a real session with a real participant. Never invent participants, times, ratings or issues.", "warn"),
        ("", None),
        ("How to use", "h"),
        ("1. Participants: one row per real participant, using the anonymous code U1, U2, … (yellow cells). Do not type names. Record consent and whether the session was recorded.", None),
        ("2. Task log: one row per participant × task is already listed (U1 to U10, T1 to T10). Fill Outcome, Time, Errors, Hints and SEQ only for tasks the participant actually attempted. Leave all other rows blank: blank rows are not counted. Extra rows (yellow, below the grid) can be used for repeated attempts or a participant above U10.", None),
        ("3. SUS: type each participant's 10 answers (1 to 5). The score (0 to 100) is calculated.", None),
        ("4. Issues: log every problem seen, with the task, the screen, a severity (High, Medium, Low), the participants affected and the requirement.", None),
        ("5. Fix log: the Issue ID, where it was found and the severity come from the Issues sheet. You add what was fixed, when, and the retest result.", None),
        ("6. Summary: calculated automatically (N, success rate, median time, mean errors, mean SEQ per task, SUS mean, pass or fail against the targets). Copy its numbers into the report.", None),
        ("7. Before-After: Round 1 numbers come from the Summary. Enter the retest numbers only if a real retest was run after the fixes.", None),
        ("", None),
        ("Rules", "h"),
        ("Outcome is Success (done unaided), Partial (done after a hint or a detour) or Fail. Time is counted for Success attempts only (stopwatch from the end of the reading to the goal screen). An error is a tap that leads away from the goal, an error or validation message the participant triggers, or going back to undo. A fault of the app or the network is not a participant error: repeat the task and say so in the notes.", None),
        ("Targets (yellow-blue cells on the Summary sheet) come from section 6.4 of the M02 report: Success at least 80 % per task, time at or under the task target, at most 2 errors per task (T1: 0 wrong turns), mean SEQ at least 5.5, mean SUS at least 68, no open High issue. T3 time target was adjusted from 120 s to 150 s (phone photo prompt and gallery); every other target is unchanged.", None),
        ("The Summary needs at least 5 participants who gave consent (brief). Tasks with fewer than 5 attempts are flagged “Few”: report them with their N.", None),
        ("Colour code: yellow = you type here; blue = target (edit only if the report changes it); grey = fixed text; no fill = formula.", None),
        ("Companion files: Gedara_App_Test_Session_Script.pdf (facilitator script and observation sheets), README_usability.md, usability3.template.json.", None),
    ]
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Arial", size=10, bold=(kind == "h"))
        if kind == "title":
            c.font = Font(name="Arial", size=14, bold=True, color=NAVY)
        if kind == "warn":
            c.font = Font(name="Arial", size=10, bold=True, color="AA3029")
    ws.column_dimensions["A"].width = 125

    # ------------------------------------------------------------ Lists (hidden helper)
    wl = wb.create_sheet("Lists")
    wl.cell(row=1, column=1, value="Screens").font = bold
    for i, s in enumerate(D.SCREENS, 2):
        wl.cell(row=i, column=1, value=s)
    wl.sheet_state = "hidden"

    # ------------------------------------------------------------ Participants
    wp = wb.create_sheet("Participants")
    header(wp, 1, ["Code", "Real background", "Roles tested", "Age group (18+ only)", "Device", "Connection", "Session date",
                   "Consent given (Yes/No)", "Recorded (Yes/No)", "Recording file name", "Tasks attempted", "Notes"], 34)
    for r in range(2, NP + 2):
        c = wp.cell(row=r, column=1, value=f"U{r-1}"); c.font = bold; c.fill = soft; c.border = box
        inputs(wp, r, 2, 10)
        inputs(wp, r, 12, 12, True)
        k = wp.cell(row=r, column=11, value=f"=COUNTIFS('Task log'!$A$2:$A${LOG_LAST},A{r},'Task log'!$C$2:$C${LOG_LAST},\"<>\")")
        k.font = base; k.border = box
        wp.cell(row=r, column=7).number_format = "dd mmm yyyy"
    for rng, lst in (("B2:B%d" % (NP + 1), D.BACKGROUNDS), ("C2:C%d" % (NP + 1), D.ROLE_OPTIONS),
                     ("D2:D%d" % (NP + 1), "18 to 24,25 to 34,35 to 44,45 to 54,55 or over"),
                     ("E2:E%d" % (NP + 1), "iPhone,Android phone,Laptop,Other"), ("F2:F%d" % (NP + 1), "Expo Go,APK,Browser"),
                     ("H2:H%d" % (NP + 1), "Yes,No"), ("I2:I%d" % (NP + 1), "Yes,No")):
        dv = DataValidation(type="list", formula1='"%s"' % lst, allow_blank=True); wp.add_data_validation(dv); dv.add(rng)
    widths(wp, [8, 26, 32, 16, 14, 13, 14, 14, 13, 30, 11, 36])
    wp.freeze_panes = "B2"
    wp.cell(row=NP + 3, column=1, value="Use only the anonymous codes. “Real background” must be honest: say if a participant is a friend or student role-playing a provider or administrator.").font = base

    # ------------------------------------------------------------ Task log
    wt = wb.create_sheet("Task log")
    header(wt, 1, ["Participant", "Task", "Outcome", "Time (s)", "Errors", "Hints given", "SEQ (1-7)", "Notes / quote"], 30)
    r = 2
    for pi in range(1, NP + 1):
        for t in D.TASK_IDS:
            for cidx, v in ((1, f"U{pi}"), (2, t)):
                c = wt.cell(row=r, column=cidx, value=v); c.font = base; c.fill = soft; c.border = box
            inputs(wt, r, 3, 8, True)
            r += 1
    for r2 in range(FIRST_FREE, FIRST_FREE + 50):
        inputs(wt, r2, 1, 8, True)
    last_row = FIRST_FREE + 49
    for rng, f1 in ((f"C2:C{last_row}", '"Success,Partial,Fail"'),):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True); wt.add_data_validation(dv); dv.add(rng)
    dvp = DataValidation(type="list", formula1="Participants!$A$2:$A$%d" % (NP + 1), allow_blank=True); wt.add_data_validation(dvp); dvp.add(f"A{FIRST_FREE}:A{last_row}")
    dvt = DataValidation(type="list", formula1='"%s"' % ",".join(D.TASK_IDS), allow_blank=True); wt.add_data_validation(dvt); dvt.add(f"B{FIRST_FREE}:B{last_row}")
    dvn = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True); wt.add_data_validation(dvn); dvn.add(f"D2:D{last_row}")
    dvw = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True); wt.add_data_validation(dvw); dvw.add(f"E2:F{last_row}")
    dvs = DataValidation(type="whole", operator="between", formula1="1", formula2="7", allow_blank=True); wt.add_data_validation(dvs); dvs.add(f"G2:G{last_row}")
    widths(wt, [12, 8, 12, 10, 9, 11, 10, 60]); wt.freeze_panes = "C2"
    wt.auto_filter.ref = f"A1:H{last_row}"
    for val, fill in (("Success", green), ("Fail", red), ("Partial", amber)):
        wt.conditional_formatting.add(f"C2:C{last_row}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill))

    # ------------------------------------------------------------ SUS
    su = wb.create_sheet("SUS")
    header(su, 1, ["Participant"] + [f"Q{i}" for i in range(1, 11)] + ["Contribution sum", "SUS score (0-100)", "Reading"], 32)
    for r in range(2, NP + 2):
        c = su.cell(row=r, column=1, value=f"U{r-1}"); c.font = bold; c.fill = soft; c.border = box
        inputs(su, r, 2, 11)
        odd = "+".join(f"({L(c)}{r}-1)" for c in (2, 4, 6, 8, 10))
        even = "+".join(f"(5-{L(c)}{r})" for c in (3, 5, 7, 9, 11))
        su.cell(row=r, column=12, value=f'=IF(COUNT(B{r}:K{r})=10,{odd}+{even},"")')
        su.cell(row=r, column=13, value=f'=IF(L{r}="","",L{r}*2.5)'); su.cell(row=r, column=13).number_format = "0.0"
        su.cell(row=r, column=14, value=f'=IF(M{r}="","",IF(M{r}>=80.3,"Excellent",IF(M{r}>=68,"Good / above average",IF(M{r}>=51,"Below average","Poor"))))')
        for cidx in (12, 13, 14):
            su.cell(row=r, column=cidx).font = base; su.cell(row=r, column=cidx).border = box
    dvq = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); su.add_data_validation(dvq); dvq.add(f"B2:K{NP+1}")
    widths(su, [12] + [6] * 10 + [14, 18, 22])
    lr = NP + 3
    su.cell(row=lr, column=1, value="Mean").font = bold
    su.cell(row=lr, column=13, value='=IFERROR(AVERAGE(M2:M%d),"")' % (NP + 1)).number_format = "0.0"
    su.cell(row=lr, column=13).font = bold
    su.cell(row=lr + 2, column=1, value="Answers: 1 = strongly disagree … 5 = strongly agree. The ten statements (System Usability Scale, Brooke 1996):").font = bold
    for i, s in enumerate(D.SUS_ITEMS, 1):
        su.cell(row=lr + 2 + i, column=1, value=f"Q{i}").font = bold
        su.cell(row=lr + 2 + i, column=2, value=s).font = base
    su.cell(row=lr + 14, column=1, value="Scoring: odd items answer minus 1; even items 5 minus answer; add the ten contributions and multiply by 2.5 (range 0 to 100). A score is shown only when all 10 answers are filled in. Average SUS is about 68 (Sauro & Lewis, 2016).").font = base
    su.freeze_panes = "B2"

    # ------------------------------------------------------------ Issues
    wi = wb.create_sheet("Issues")
    header(wi, 1, ["ID", "Task", "Interface (screen)", "Issue: what happened", "Severity", "Participants affected (number)",
                   "Participant codes", "Requirement", "Status", "Found in", "Notes"], 42)
    for k in range(1, NI + 1):
        r = k + 1
        c = wi.cell(row=r, column=1, value=f"UT-{k:02d}"); c.font = bold; c.fill = soft; c.border = box
        inputs(wi, r, 2, 11, True)
    for rng, f1 in ((f"B2:B{NI+1}", '"%s"' % ",".join(D.TASK_IDS)), (f"E2:E{NI+1}", '"High,Medium,Low"'),
                    (f"I2:I{NI+1}", '"Open,Fixed,Won\'t fix"'), (f"J2:J{NI+1}", '"Round 1,Retest"'),
                    (f"C2:C{NI+1}", "Lists!$A$2:$A$%d" % (len(D.SCREENS) + 1))):
        dv = DataValidation(type="list", formula1=f1, allow_blank=True, showErrorMessage=(rng[0] != "C"))
        wi.add_data_validation(dv); dv.add(rng)
    widths(wi, [8, 7, 24, 58, 11, 16, 16, 14, 11, 11, 40])
    for sev, fill in (("High", red), ("Medium", amber), ("Low", PatternFill("solid", bgColor="E3EEF9"))):
        wi.conditional_formatting.add(f"E2:E{NI+1}", CellIsRule(operator="equal", formula=[f'"{sev}"'], fill=fill))
    wi.freeze_panes = "B2"

    # ------------------------------------------------------------ Fix log
    wf = wb.create_sheet("Fix log")
    header(wf, 1, ["Issue ID", "Found in", "Severity", "Fix made (what changed in the app)", "Date fixed", "Retest result", "Retest notes"], 32)
    for k in range(1, NI + 1):
        r = k + 1
        wf.cell(row=r, column=1, value=f'=IF(Issues!D{r}="","",Issues!A{r})')
        wf.cell(row=r, column=2, value=f'=IF(Issues!D{r}="","",Issues!B{r}&IF(Issues!C{r}<>""," · "&Issues!C{r},"")&IF(Issues!J{r}<>""," · "&Issues!J{r},""))')
        wf.cell(row=r, column=3, value=f'=IF(Issues!D{r}="","",Issues!E{r})')
        for cidx in (1, 2, 3):
            c = wf.cell(row=r, column=cidx); c.font = base; c.border = box; c.alignment = wrap
        inputs(wf, r, 4, 7, True)
        wf.cell(row=r, column=5).number_format = "dd mmm yyyy"
    dvr = DataValidation(type="list", formula1='"Retested with a user: fixed,Retested with a user: still a problem,Checked by automated test only,Not retested"', allow_blank=True)
    wf.add_data_validation(dvr); dvr.add(f"F2:F{NI+1}")
    widths(wf, [10, 34, 11, 56, 13, 34, 40])
    for sev, fill in (("High", red), ("Medium", amber)):
        wf.conditional_formatting.add(f"C2:C{NI+1}", CellIsRule(operator="equal", formula=[f'"{sev}"'], fill=fill))
    wf.freeze_panes = "A2"
    wf.cell(row=NI + 3, column=1, value="Rows appear here when a row on the Issues sheet has an issue text. “Checked by automated test only” means the app's end-to-end tests pass but no participant saw the fix: say so in the report.").font = base

    # ------------------------------------------------------------ Summary
    sm = wb.create_sheet("Summary")
    header(sm, 1, ["Task", "Role", "Task name", "N (attempts)", "Success", "Partial", "Fail", "Success rate", "Median time, Success only (s)",
                   "Mean errors", "Mean SEQ", "Target success", "Target time (s)", "Target max mean errors", "Target SEQ",
                   "Success OK?", "Time OK?", "Errors OK?", "SEQ OK?", "Task result", "Attempts ≥ 5?"], 58)
    R = f"'Task log'!$B$2:$B${LOG_LAST}"; O = f"'Task log'!$C$2:$C${LOG_LAST}"; T_ = f"'Task log'!$D$2:$D${LOG_LAST}"
    E_ = f"'Task log'!$E$2:$E${LOG_LAST}"; Q = f"'Task log'!$G$2:$G${LOG_LAST}"
    for i, t in enumerate(D.TASKS, 2):
        sm.cell(row=i, column=1, value=t["id"]); sm.cell(row=i, column=2, value=t["role"]); sm.cell(row=i, column=3, value=t["title"])
        sm.cell(row=i, column=4, value=f'=COUNTIFS({R},A{i},{O},"<>")')
        sm.cell(row=i, column=5, value=f'=COUNTIFS({R},A{i},{O},"Success")')
        sm.cell(row=i, column=6, value=f'=COUNTIFS({R},A{i},{O},"Partial")')
        sm.cell(row=i, column=7, value=f'=COUNTIFS({R},A{i},{O},"Fail")')
        sm.cell(row=i, column=8, value=f'=IF(D{i}=0,"",E{i}/D{i})'); sm.cell(row=i, column=8).number_format = "0%"
        sm.cell(row=i, column=9).value = ArrayFormula(f"I{i}", f'=IFERROR(MEDIAN(IF(({R}=A{i})*({O}="Success")*({T_}<>""),{T_})),"")')
        sm.cell(row=i, column=9).number_format = "0"
        sm.cell(row=i, column=10, value=f'=IFERROR(AVERAGEIFS({E_},{R},A{i},{O},"<>"),"")'); sm.cell(row=i, column=10).number_format = "0.0"
        sm.cell(row=i, column=11, value=f'=IFERROR(AVERAGEIFS({Q},{R},A{i},{O},"<>"),"")'); sm.cell(row=i, column=11).number_format = "0.0"
        sm.cell(row=i, column=12, value=D.TARGET_SUCCESS); sm.cell(row=i, column=12).number_format = "0%"
        sm.cell(row=i, column=13, value=t["time"]); sm.cell(row=i, column=14, value=t["errors"]); sm.cell(row=i, column=15, value=D.TARGET_SEQ)
        sm.cell(row=i, column=16, value=f'=IF(D{i}=0,"",IF(H{i}>=L{i},"Yes","No"))')
        sm.cell(row=i, column=17, value=f'=IF(I{i}="","",IF(I{i}<=M{i},"Yes","No"))')
        sm.cell(row=i, column=18, value=f'=IF(J{i}="","",IF(J{i}<=N{i},"Yes","No"))')
        sm.cell(row=i, column=19, value=f'=IF(K{i}="","",IF(K{i}>=O{i},"Yes","No"))')
        sm.cell(row=i, column=20, value=f'=IF(D{i}=0,"",IF(COUNTIF(P{i}:S{i},"No")>0,"Fail",IF(COUNTIF(P{i}:S{i},"Yes")=4,"Pass","Incomplete")))')
        sm.cell(row=i, column=21, value=f'=IF(D{i}=0,"",IF(D{i}>=5,"Yes","Few"))')
        for cidx in range(1, 22):
            c = sm.cell(row=i, column=cidx); c.border = box
            c.font = bold if cidx in (1, 20) else base
            if cidx == 3:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if cidx in (12, 13, 14, 15):
                c.fill = tgt_fill
            elif i % 2 == 1:
                c.fill = soft
    last = len(D.TASKS) + 1
    for col in ("P", "Q", "R", "S"):
        rng = f"{col}2:{col}{last}"
        sm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Yes"'], fill=green))
        sm.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"No"'], fill=red))
    sm.conditional_formatting.add(f"T2:T{last}", CellIsRule(operator="equal", formula=['"Pass"'], fill=green))
    sm.conditional_formatting.add(f"T2:T{last}", CellIsRule(operator="equal", formula=['"Fail"'], fill=red))
    sm.conditional_formatting.add(f"U2:U{last}", CellIsRule(operator="equal", formula=['"Few"'], fill=amber))
    sm.cell(row=last + 1, column=1, value="Targets (blue cells): M02 report section 6.4. T3 time ADJUSTED to 150 s (phone photo prompt and gallery). T1 allows 0 errors (0 wrong turns); other tasks at most 2.").font = base

    r0 = last + 3
    sm.cell(row=r0, column=1, value="Participants").font = bold
    sm.cell(row=r0 + 1, column=1, value="Participants who gave consent (N)").font = base
    sm.cell(row=r0 + 1, column=4, value=f'=SUMPRODUCT((Participants!A2:A{NP+1}<>"")*(Participants!H2:H{NP+1}="Yes"))').font = bold
    sm.cell(row=r0 + 2, column=1, value="Minimum required by the brief: 5").font = base
    sm.cell(row=r0 + 2, column=4, value=f'=IF(D{r0+1}>=5,"Met","Not yet")').font = bold
    sm.cell(row=r0 + 3, column=1, value="Participants with at least one task logged").font = base
    sm.cell(row=r0 + 3, column=4, value=f'=SUMPRODUCT(--(Participants!K2:K{NP+1}>0))').font = bold
    sm.cell(row=r0 + 4, column=1, value="Sessions recorded (with consent)").font = base
    sm.cell(row=r0 + 4, column=4, value=f'=COUNTIFS(Participants!I2:I{NP+1},"Yes",Participants!H2:H{NP+1},"Yes")').font = bold

    r1 = r0 + 6
    sm.cell(row=r1, column=1, value="System Usability Scale").font = bold
    sm.cell(row=r1 + 1, column=1, value="Mean SUS score").font = base
    sm.cell(row=r1 + 1, column=4, value=f'=IFERROR(AVERAGE(SUS!M2:M{NP+1}),"")').number_format = "0.0"
    sm.cell(row=r1 + 1, column=4).font = bold
    sm.cell(row=r1 + 2, column=1, value="Questionnaires with a score (n)").font = base
    sm.cell(row=r1 + 2, column=4, value=f'=COUNT(SUS!M2:M{NP+1})').font = bold
    sm.cell(row=r1 + 3, column=1, value="Target: %d (average)" % D.TARGET_SUS).font = base
    sm.cell(row=r1 + 3, column=4, value=f'=IF(D{r1+1}="","",IF(D{r1+1}>={D.TARGET_SUS},"Met","Below target"))').font = bold

    r2 = r1 + 5
    sm.cell(row=r2, column=1, value="Issues found in the user sessions").font = bold
    sm.cell(row=r2, column=3, value="Logged").font = bold; sm.cell(row=r2, column=4, value="Still open").font = bold
    sm.cell(row=r2, column=5, value="Fixed").font = bold
    for k, sev in enumerate(["High", "Medium", "Low"]):
        rr = r2 + 1 + k
        sm.cell(row=rr, column=1, value=sev).font = base
        sm.cell(row=rr, column=3, value=f'=COUNTIFS(Issues!$E$2:$E${NI+1},"{sev}")').font = base
        sm.cell(row=rr, column=4, value=f'=COUNTIFS(Issues!$E$2:$E${NI+1},"{sev}",Issues!$I$2:$I${NI+1},"Open")').font = base
        sm.cell(row=rr, column=5, value=f'=COUNTIFS(Issues!$E$2:$E${NI+1},"{sev}",Issues!$I$2:$I${NI+1},"Fixed")').font = base
    rh = r2 + 4
    sm.cell(row=rh, column=1, value="No open High issue (M02 threshold)").font = base
    sm.cell(row=rh, column=4, value=f'=IF(C{r2+1}=0,"None logged",IF(D{r2+1}=0,"Met","Not met"))').font = bold

    r3 = rh + 2
    sm.cell(row=r3, column=1, value="Task results").font = bold
    sm.cell(row=r3 + 1, column=1, value="Tasks with a Pass result").font = base
    sm.cell(row=r3 + 1, column=4, value=f'=COUNTIF(T2:T{last},"Pass")').font = bold
    sm.cell(row=r3 + 2, column=1, value="Tasks attempted at least once").font = base
    sm.cell(row=r3 + 2, column=4, value=f'=SUMPRODUCT(--(D2:D{last}>0))').font = bold
    widths(sm, [30, 13, 34, 10, 9, 9, 7, 10, 15, 10, 9, 10, 10, 12, 9, 9, 8, 9, 8, 12, 11])
    sm.freeze_panes = "D2"

    # ------------------------------------------------------------ Before-After
    ba = wb.create_sheet("Before-After")
    ba.cell(row=1, column=1, value="Before / after: Round 1 against the retest").font = Font(name="Arial", size=14, bold=True, color=NAVY)
    notes = [
        "Before = Round 1: all participants, the app exactly as first tested. These columns are calculated from the Summary sheet.",
        "After = the retest: the same tasks, run again after the fixes in the Fix log, on the screens that changed. Type these numbers only from real retest sessions (new participants or the same ones: say which in the report).",
        "If no retest with participants was run, leave the After columns empty. Do not estimate them. In the report, state which fixes were checked only by the automated end-to-end tests (Fix log: “Checked by automated test only”).",
        "Change columns are calculated only when both Before and After exist. A positive change in success rate or SEQ and a negative change in time or errors are improvements.",
    ]
    for i, t in enumerate(notes, 3):
        c = ba.cell(row=i, column=1, value=t); c.font = base; c.alignment = Alignment(wrap_text=True, vertical="top")
        ba.merge_cells(start_row=i, start_column=1, end_row=i, end_column=16)
        ba.row_dimensions[i].height = 28
    hr = 8
    header(ba, hr, ["Task", "Task name", "Before: N", "Before: success rate", "Before: median time (s)", "Before: mean errors", "Before: mean SEQ",
                    "After: N (retest)", "After: success rate", "After: median time (s)", "After: mean errors", "After: mean SEQ",
                    "Change: success rate", "Change: time (s)", "Change: errors", "Change: SEQ"], 46)
    for k, t in enumerate(D.TASKS):
        r = hr + 1 + k; s = k + 2
        ba.cell(row=r, column=1, value=t["id"]); ba.cell(row=r, column=2, value=t["title"])
        ba.cell(row=r, column=3, value=f"=Summary!D{s}")
        ba.cell(row=r, column=4, value=f"=Summary!H{s}"); ba.cell(row=r, column=4).number_format = "0%"
        ba.cell(row=r, column=5, value=f"=Summary!I{s}"); ba.cell(row=r, column=5).number_format = "0"
        ba.cell(row=r, column=6, value=f"=Summary!J{s}"); ba.cell(row=r, column=6).number_format = "0.0"
        ba.cell(row=r, column=7, value=f"=Summary!K{s}"); ba.cell(row=r, column=7).number_format = "0.0"
        for cidx in range(8, 13):
            c = ba.cell(row=r, column=cidx); c.fill = inp
        ba.cell(row=r, column=9).number_format = "0%"
        for cidx, (b_, a_, fmt) in zip((13, 14, 15, 16), (("D", "I", "0%"), ("E", "J", "0"), ("F", "K", "0.0"), ("G", "L", "0.0"))):
            ba.cell(row=r, column=cidx, value=f'=IF(OR({b_}{r}="",{a_}{r}=""),"",{a_}{r}-{b_}{r})').number_format = fmt
        for cidx in range(1, 17):
            c = ba.cell(row=r, column=cidx); c.border = box
            if c.font != bold:
                c.font = base
            if cidx == 2:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    dvp2 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True); ba.add_data_validation(dvp2)
    dvp2.add(f"I{hr+1}:I{hr+len(D.TASKS)}")
    rs = hr + len(D.TASKS) + 2
    ba.cell(row=rs, column=1, value="SUS (mean)").font = bold
    ba.cell(row=rs, column=4, value="Before").font = bold; ba.cell(row=rs, column=9, value="After (retest)").font = bold; ba.cell(row=rs, column=13, value="Change").font = bold
    ba.cell(row=rs + 1, column=4, value="=Summary!D%d" % (r1 + 1)).number_format = "0.0"
    c = ba.cell(row=rs + 1, column=9); c.fill = inp; c.border = box; c.number_format = "0.0"
    ba.cell(row=rs + 1, column=13, value=f'=IF(OR(D{rs+1}="",I{rs+1}=""),"",I{rs+1}-D{rs+1})').number_format = "0.0"
    for cidx in (4, 13):
        ba.cell(row=rs + 1, column=cidx).font = bold
    widths(ba, [8, 34, 9, 12, 13, 12, 11, 11, 12, 13, 12, 11, 12, 11, 10, 10])
    ba.freeze_panes = f"C{hr+1}"

    order = ["Read me", "Participants", "Task log", "SUS", "Issues", "Fix log", "Summary", "Before-After", "Lists"]
    wb._sheets = [wb[n] for n in order]
    wb.calculation.fullCalcOnLoad = True
    if hook:
        hook(wb)
    wb.save(out)
    return wb


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "IT3060_Milestone03_Usability_Workbook.xlsx"
    build(out)
    print("saved", out)
